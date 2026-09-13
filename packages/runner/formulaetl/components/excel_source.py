"""Excel Source — read .xlsx/.xls sheets into row dicts."""

from __future__ import annotations

from typing import Any

from formulaetl.sdk.base import BaseComponent
from formulaetl.sdk.context import ComponentResult, Metrics, RunContext, timed
from formulaetl.sdk.registry import register


def _parse_sheet_name(raw: Any) -> str | int:
    if raw is None or raw == "":
        return 0
    if isinstance(raw, int):
        return raw
    s = str(raw).strip()
    if s.isdigit() or (s.startswith("-") and s[1:].isdigit()):
        return int(s)
    return s


def _parse_range(raw: Any) -> str | None:
    if raw is None or raw == "":
        return None
    return str(raw).strip() or None


def _normalize_value(v: Any) -> Any:
    if v is None:
        return None
    try:
        import math

        if isinstance(v, float) and math.isnan(v):
            return None
    except Exception:
        pass
    try:
        import numpy as np

        if isinstance(v, (np.integer,)):
            return int(v)
        if isinstance(v, (np.floating,)):
            if np.isnan(v):
                return None
            return float(v)
        if isinstance(v, (np.bool_,)):
            return bool(v)
    except Exception:
        pass
    # pandas Timestamp → ISO date string when time is midnight
    try:
        import pandas as pd

        if isinstance(v, pd.Timestamp):
            if v.hour == 0 and v.minute == 0 and v.second == 0:
                return v.strftime("%Y-%m-%d")
            return v.isoformat()
    except Exception:
        pass
    return v


def _read_with_openpyxl_range(
    path: Any, sheet: str | int, cell_range: str, has_header: bool
) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if isinstance(sheet, int):
            ws = wb.worksheets[sheet]
        else:
            ws = wb[str(sheet)]
        cells = ws[cell_range]
        if not isinstance(cells, tuple):
            cells = ((cells,),)
        elif cells and not isinstance(cells[0], tuple):
            cells = (cells,)
        matrix = [[_normalize_value(c.value) for c in row] for row in cells]
    finally:
        wb.close()

    if not matrix:
        return []
    if has_header:
        cols: list[str] = []
        seen: dict[str, int] = {}
        for i, c in enumerate(matrix[0]):
            name = str(c) if c is not None else f"col_{i}"
            if name in seen:
                seen[name] += 1
                name = f"{name}_{seen[name]}"
            else:
                seen[name] = 0
            cols.append(name)
        return [dict(zip(cols, row)) for row in matrix[1:]]
    cols = [f"col_{i}" for i in range(len(matrix[0]))]
    return [dict(zip(cols, row)) for row in matrix]


@register
class ExcelSource(BaseComponent):
    component_type = "excel_source"
    display_name = "Excel Source"
    category = "source"
    config_schema = {
        "type": "object",
        "required": ["path"],
        "properties": {
            "path": {"type": "string", "description": "Path to .xlsx/.xls relative to work_dir"},
            "sheet_name": {
                "type": ["string", "integer"],
                "description": "Sheet name or 0-based index (default 0)",
            },
            "has_header": {"type": "boolean", "default": True},
            "range": {
                "type": "string",
                "description": "Optional cell range, e.g. A1:G100",
            },
        },
    }
    parameters = [
        {
            "key": "path",
            "label": "Excel path",
            "type": "string",
            "required": True,
            "help": "Path to .xlsx/.xls relative to work_dir",
        },
        {
            "key": "sheet_name",
            "label": "Sheet",
            "type": "string",
            "required": False,
            "default": "0",
            "help": "Sheet name or 0-based index",
        },
        {
            "key": "has_header",
            "label": "Has header",
            "type": "boolean",
            "required": False,
            "default": True,
            "help": "First row contains column names",
        },
        {
            "key": "range",
            "label": "Range",
            "type": "string",
            "required": False,
            "help": "Optional cell range (e.g. A1:G100)",
        },
    ]

    def run(self, ctx: RunContext, rows: list[dict[str, Any]] | None = None) -> ComponentResult:
        metrics = Metrics()
        with timed(metrics):
            import pandas as pd

            path = ctx.resolve(self.config["path"])
            if not path.exists():
                raise FileNotFoundError(f"ExcelSource: file not found: {path}")

            suffix = path.suffix.lower()
            if suffix not in (".xlsx", ".xlsm", ".xltx", ".xltm", ".xls"):
                raise ValueError(
                    f"ExcelSource: unsupported extension {suffix!r} (use .xlsx/.xls)"
                )

            sheet = _parse_sheet_name(self.config.get("sheet_name", 0))
            has_header = self.config.get("has_header", True)
            if isinstance(has_header, str):
                has_header = has_header.strip().lower() in ("1", "true", "yes", "y")
            cell_range = _parse_range(self.config.get("range"))

            if cell_range and suffix != ".xls":
                out_rows = _read_with_openpyxl_range(path, sheet, cell_range, has_header)
            else:
                header = 0 if has_header else None
                engine = "openpyxl" if suffix != ".xls" else None
                read_kwargs: dict[str, Any] = {
                    "sheet_name": sheet,
                    "header": header,
                    "dtype": object,
                }
                if engine:
                    read_kwargs["engine"] = engine
                df = pd.read_excel(path, **read_kwargs)

                cols: list[str] = []
                seen: dict[str, int] = {}
                for i, c in enumerate(df.columns):
                    name = str(c) if c is not None and str(c) != "nan" else f"col_{i}"
                    if name in seen:
                        seen[name] += 1
                        name = f"{name}_{seen[name]}"
                    else:
                        seen[name] = 0
                    cols.append(name)
                df.columns = cols

                out_rows = []
                for rec in df.to_dict(orient="records"):
                    out_rows.append({k: _normalize_value(v) for k, v in rec.items()})

            metrics.rows_in = 0
            metrics.rows_out = len(out_rows)
            ctx.emit(
                f"ExcelSource: read {path.name} sheet={sheet!r} → {len(out_rows)} rows"
            )

        return ComponentResult(
            rows=out_rows,
            metrics=metrics,
            artifacts={"path": str(path), "sheet_name": sheet},
            side_effects={"path": str(path), "sheet_name": sheet, "rows": len(out_rows)},
        )
